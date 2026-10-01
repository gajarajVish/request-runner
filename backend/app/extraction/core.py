from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime
from typing import Any, Callable

from PIL import Image

MAX_RENDER_PAGES = 8
LOW_TEXT_CHARS = 40
MAX_SHEET_ROWS = 400
MAX_SHEET_COLS = 30
_SIGNATURE_HINT = re.compile(
    r"signature|signed|/s/|\bby:\s|sign here|authorized signatory|_{8,}\s*\n\s*name:", re.IGNORECASE
)


class Unreadable(Exception):
    """The file can't be assessed. The message is shown to the requester and provider."""


PutBlob = Callable[[bytes], str]


def extract(data: bytes, detected: str, filename: str, put_blob: PutBlob) -> dict[str, Any]:
    if detected == "pdf":
        return _pdf(data, put_blob)
    if detected in ("xlsx", "xls", "csv"):
        return _sheet(data, detected)
    if detected == "docx":
        return _docx(data)
    if detected in ("png", "jpeg", "heic"):
        return _image(data, detected, put_blob)
    if detected == "encrypted_office":
        raise Unreadable("the Office file is password-protected")
    raise Unreadable(f"unsupported type {detected}")


# --------------------------------------------------------------------------- PDF


def _pdf(data: bytes, put_blob: PutBlob) -> dict[str, Any]:
    import pypdfium2 as pdfium

    try:
        pdf = pdfium.PdfDocument(data)
    except pdfium.PdfiumError as e:
        msg = str(e).lower()
        if "password" in msg:
            raise Unreadable("the PDF is password-protected") from e
        raise Unreadable("the PDF is corrupt or not a valid PDF") from e
    n = len(pdf)
    if n == 0:
        raise Unreadable("the PDF has no pages")
    segments, images = [], []
    rendered = 0
    blank_pages = 0
    for i in range(n):
        page = pdf[i]
        tp = page.get_textpage()
        text = tp.get_text_range() or ""
        text = text.replace("\r\n", "\n").replace("\r", "\n")
        loc = {"type": "pdf", "page": i + 1}
        if text.strip():
            segments.append({"loc": loc, "label": f"p.{i + 1}", "text": text.strip()})
        why = None
        if len(text.strip()) < LOW_TEXT_CHARS:
            why = "little or no extractable text (scan or image)"
        elif _SIGNATURE_HINT.search(text):
            why = "signature block - check visually"
        if why and rendered < MAX_RENDER_PAGES:
            img = page.render(scale=1.6).to_pil().convert("RGB")
            if _is_blank(img):
                blank_pages += 1
                if not text.strip():
                    continue
            buf = io.BytesIO()
            img.save(buf, "PNG", optimize=True)
            images.append(
                {"loc": loc, "label": f"p.{i + 1} (image)", "blob": put_blob(buf.getvalue()), "media_type": "image/png", "why": why}
            )
            rendered += 1
    if not segments and not images:
        raise Unreadable("the PDF pages are blank")
    return {"segments": segments, "images": images, "meta": {"pages": n, "blank_pages": blank_pages}}


def _is_blank(img: Image.Image) -> bool:
    gray = img.convert("L").resize((64, 64))
    lo, hi = gray.getextrema()
    return hi - lo < 12


# --------------------------------------------------------------------------- spreadsheets


def _fmt(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == datetime.min.time() else v.isoformat(sep=" ")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float):
        # spreadsheet floats carry binary noise (346037.7600000001); show what the cell displays
        return str(int(v)) if v.is_integer() else format(v, ".12g")
    return str(v).strip()


def _col(n: int) -> str:
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def _sheet_rows(data: bytes, detected: str) -> list[tuple[str, list[list[Any]]]]:
    if detected == "xlsx":
        import openpyxl

        try:
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as e:  # noqa: BLE001
            raise Unreadable("the spreadsheet is corrupt or password-protected") from e
        out = []
        for ws in wb.worksheets:
            rows = [list(r) for r in ws.iter_rows(max_row=MAX_SHEET_ROWS, max_col=MAX_SHEET_COLS, values_only=True)]
            out.append((ws.title, rows))
        return out
    if detected == "xls":
        import xlrd

        try:
            book = xlrd.open_workbook(file_contents=data)
        except Exception as e:  # noqa: BLE001
            raise Unreadable("the .xls file is corrupt or password-protected") from e
        out = []
        for sh in book.sheets():
            rows = []
            for r in range(min(sh.nrows, MAX_SHEET_ROWS)):
                row = []
                for c in range(min(sh.ncols, MAX_SHEET_COLS)):
                    cell = sh.cell(r, c)
                    v = cell.value
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        v = xlrd.xldate.xldate_as_datetime(v, book.datemode)
                    row.append(v)
                rows.append(row)
            out.append((sh.name, rows))
        return out
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
    rows = [r[:MAX_SHEET_COLS] for r in list(csv.reader(io.StringIO(text)))[:MAX_SHEET_ROWS]]
    return [("csv", rows)]


def _sheet(data: bytes, detected: str) -> dict[str, Any]:
    segments = []
    sheets_meta = []
    for title, rows in _sheet_rows(data, detected):
        lines = []
        cells: dict[str, str] = {}
        last_row = last_col = 0
        for r_i, row in enumerate(rows, start=1):
            vals = []
            for c_i, v in enumerate(row, start=1):
                s = _fmt(v)
                if s:
                    coord = f"{_col(c_i)}{r_i}"
                    cells[coord] = s
                    vals.append(f"{coord}={s}")
                    last_row, last_col = max(last_row, r_i), max(last_col, c_i)
            if vals:
                lines.append(" | ".join(vals))
        sheets_meta.append({"sheet": title, "cells": len(cells)})
        if not cells:
            continue
        rng = f"A1:{_col(last_col)}{last_row}"
        segments.append(
            {"loc": {"type": "sheet", "sheet": title, "range": rng}, "label": f"sheet '{title}' {rng}", "text": "\n".join(lines), "cells": cells}
        )
    if not segments:
        raise Unreadable("the spreadsheet is empty")
    return {"segments": segments, "images": [], "meta": {"sheets": sheets_meta}}


# --------------------------------------------------------------------------- DOCX


def _docx(data: bytes) -> dict[str, Any]:
    import docx

    try:
        d = docx.Document(io.BytesIO(data))
    except Exception as e:  # noqa: BLE001
        raise Unreadable("the Word document is corrupt or password-protected") from e
    segments = []
    for i, p in enumerate(d.paragraphs, start=1):
        if p.text.strip():
            segments.append({"loc": {"type": "docx", "paragraph": i}, "label": f"paragraph {i}", "text": p.text.strip()})
    for t_i, table in enumerate(d.tables, start=1):
        for r_i, row in enumerate(table.rows, start=1):
            text = " | ".join(c.text.strip() for c in row.cells)
            if text.strip(" |"):
                segments.append(
                    {"loc": {"type": "docx", "table": t_i, "row": r_i}, "label": f"table {t_i} row {r_i}", "text": text}
                )
    if not segments:
        raise Unreadable("the Word document is empty")
    return {"segments": segments, "images": [], "meta": {"paragraphs": len(d.paragraphs), "tables": len(d.tables)}}


# --------------------------------------------------------------------------- images


def _image(data: bytes, detected: str, put_blob: PutBlob) -> dict[str, Any]:
    if detected == "heic":
        try:
            import pillow_heif

            pillow_heif.register_heif_opener()
        except Exception as e:  # noqa: BLE001
            raise Unreadable("HEIC images are not supported on this server") from e
    try:
        img = Image.open(io.BytesIO(data))
        img.load()
    except Exception as e:  # noqa: BLE001
        raise Unreadable("the image is corrupt") from e
    if img.width < 200 or img.height < 200:
        raise Unreadable(f"the image is too small to read ({img.width}x{img.height}px)")
    img = img.convert("RGB")
    if max(img.size) > 2000:
        img.thumbnail((2000, 2000))
    if _is_blank(img):
        raise Unreadable("the image is blank")
    buf = io.BytesIO()
    img.save(buf, "PNG", optimize=True)
    return {
        "segments": [],
        "images": [{"loc": {"type": "image"}, "label": "image", "blob": put_blob(buf.getvalue()), "media_type": "image/png", "why": "image file"}],
        "meta": {"width": img.width, "height": img.height, "original_type": detected},
    }


# --------------------------------------------------------------------------- injection heuristics

_INJECTION = re.compile(
    r"(ignore (all |any )?(previous|prior|above) (instructions|messages)"
    r"|disregard (the|all|your) (instructions|checklist|rules)"
    r"|mark (this|the|all)( \w+){0,3} (as )?(complete|completed|met|satisfied|done)"
    r"|you are (an? )?(ai|assistant|agent|language model)"
    r"|system prompt"
    r"|(note|instruction)s? (to|for) (the )?(ai|agent|assistant|reviewer|bot)"
    r"|treat (this|all) (items?|requests?) as)",
    re.IGNORECASE,
)


_INJECTION_DOC_ONLY = re.compile(r"(close|complete|approve) (this|the) request", re.IGNORECASE)


def find_injection_like(text: str, *, in_document: bool = True) -> list[str]:
    """Return short snippets around instruction-like phrases aimed at the agent.

    `in_document`: inside a file, "close this request" is suspicious; in the provider's own
    email it is a legitimate way to close, so only the general patterns apply there.
    """
    hits = []
    pats = [_INJECTION, _INJECTION_DOC_ONLY] if in_document else [_INJECTION]
    for m in (m for p in pats for m in p.finditer(text or "")):
        start = max(0, m.start() - 60)
        hits.append(text[start : m.end() + 80].strip().replace("\n", " "))
    return hits[:5]
